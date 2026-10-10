"""Tests for the DOCX parser."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.parsing.docx_parser import parse_docx
from app.parsing.exceptions import DocumentParseError
from app.parsing.models import ParsedBlockType


def save(doc: DocxDocument, path: Path) -> Path:
    doc.save(str(path))
    return path


def paragraphs_of(document) -> list[str]:
    return [block.text for block in document.blocks]


# --- 1-3: basic text, Unicode, Greek ----------------------------------------


def test_parses_simple_paragraph(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Hello world.")
    parsed = parse_docx(save(doc, tmp_path / "simple.docx"))

    assert parsed.extension == ".docx"
    assert parsed.filename == "simple.docx"
    assert paragraphs_of(parsed) == ["Hello world."]
    assert parsed.blocks[0].block_type == ParsedBlockType.PARAGRAPH
    assert "Hello world." in parsed.raw_text


def test_preserves_unicode(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Café — naïve — 日本語 ✓")
    parsed = parse_docx(save(doc, tmp_path / "unicode.docx"))
    assert paragraphs_of(parsed) == ["Café — naïve — 日本語 ✓"]


def test_preserves_greek_text(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Ανάλυση δεδομένων: μέθοδοι και αποτελέσματα")
    parsed = parse_docx(save(doc, tmp_path / "greek.docx"))
    assert paragraphs_of(parsed) == ["Ανάλυση δεδομένων: μέθοδοι και αποτελέσματα"]


# --- 4-5: title handling ----------------------------------------------------


def test_detects_title_style(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("My Thesis Title", style="Title")
    doc.add_paragraph("Body text.")
    parsed = parse_docx(save(doc, tmp_path / "title.docx"))

    assert parsed.title == "My Thesis Title"
    assert parsed.blocks[0].block_type == ParsedBlockType.HEADING
    assert parsed.blocks[0].heading_level == 0


def test_does_not_treat_first_paragraph_as_title(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Just an ordinary opening line.")
    parsed = parse_docx(save(doc, tmp_path / "no_title.docx"))
    assert parsed.title is None


# --- 6-9: headings ----------------------------------------------------------


def test_detects_heading_one(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("Chapter 4", level=1)
    parsed = parse_docx(save(doc, tmp_path / "h1.docx"))
    block = parsed.blocks[0]
    assert block.block_type == ParsedBlockType.HEADING
    assert block.heading_level == 1
    assert block.heading_path == ["Chapter 4"]


def test_detects_heading_two(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("Chapter 4", level=1)
    doc.add_heading("Development of VR", level=2)
    parsed = parse_docx(save(doc, tmp_path / "h2.docx"))
    second = parsed.blocks[1]
    assert second.block_type == ParsedBlockType.HEADING
    assert second.heading_level == 2
    assert second.heading_path == ["Chapter 4", "Development of VR"]


def test_maintains_nested_heading_path(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("Chapter 4", level=1)
    doc.add_heading("Development of VR", level=2)
    doc.add_heading("Early immersive systems", level=3)
    doc.add_paragraph("Body text.")
    parsed = parse_docx(save(doc, tmp_path / "nested.docx"))

    body = parsed.blocks[-1]
    assert body.text == "Body text."
    assert body.heading_path == [
        "Chapter 4",
        "Development of VR",
        "Early immersive systems",
    ]


def test_clears_deeper_heading_levels(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("Chapter 4", level=1)
    doc.add_heading("Development of VR", level=2)
    doc.add_heading("Early immersive systems", level=3)
    doc.add_heading("Evidence of educational use", level=2)
    doc.add_paragraph("Following paragraph.")
    parsed = parse_docx(save(doc, tmp_path / "clear.docx"))

    body = parsed.blocks[-1]
    assert body.heading_path == ["Chapter 4", "Evidence of educational use"]


# --- 10: paragraph order ----------------------------------------------------


def test_preserves_multiple_paragraphs_in_order(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("First.")
    doc.add_paragraph("Second.")
    doc.add_paragraph("Third.")
    parsed = parse_docx(save(doc, tmp_path / "order.docx"))
    assert paragraphs_of(parsed) == ["First.", "Second.", "Third."]


# --- 11-13: tables ----------------------------------------------------------


def test_preserves_table_rows_and_cells(tmp_path: Path) -> None:
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=3)
    for column, value in enumerate(["A", "B", "C"]):
        table.cell(0, column).text = value
    for column, value in enumerate(["1", "2", "3"]):
        table.cell(1, column).text = value
    parsed = parse_docx(save(doc, tmp_path / "table.docx"))

    assert len(parsed.tables) == 1
    parsed_table = parsed.tables[0]
    assert parsed_table.table_index == 0
    assert parsed_table.source_locator == "table:0"
    assert [[cell.text for cell in row.cells] for row in parsed_table.rows] == [
        ["A", "B", "C"],
        ["1", "2", "3"],
    ]


def test_associates_table_with_current_heading_path(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("Results", level=1)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "value"
    parsed = parse_docx(save(doc, tmp_path / "table_heading.docx"))
    assert parsed.tables[0].heading_path == ["Results"]


def test_preserves_paragraph_table_interleaving(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Before table.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell A"
    doc.add_paragraph("After table.")
    parsed = parse_docx(save(doc, tmp_path / "interleave.docx"))

    assert paragraphs_of(parsed) == ["Before table.", "After table."]
    raw = parsed.raw_text
    assert raw.index("Before table.") < raw.index("Cell A") < raw.index("After table.")


# --- 14: empty paragraphs ---------------------------------------------------


def test_empty_paragraphs_do_not_create_blocks(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    doc.add_paragraph("Real content.")
    parsed = parse_docx(save(doc, tmp_path / "empty.docx"))
    assert paragraphs_of(parsed) == ["Real content."]


# --- 15 + errors ------------------------------------------------------------


def test_corrupt_docx_gives_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is definitely not a docx package")
    with pytest.raises(DocumentParseError):
        parse_docx(path)


def test_missing_file_gives_clear_error(tmp_path: Path) -> None:
    with pytest.raises(DocumentParseError):
        parse_docx(tmp_path / "missing.docx")


def test_file_hash_is_passed_through(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Body.")
    parsed = parse_docx(save(doc, tmp_path / "hash.docx"), file_hash="deadbeef")
    assert parsed.file_hash == "deadbeef"
