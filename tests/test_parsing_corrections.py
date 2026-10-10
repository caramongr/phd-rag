"""Tests for RAG-003 corrections: heading provenance, DOCX heading inference,
and the canonical ordered ``content_items`` representation.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pydantic import ValidationError

from app.parsing.docx_parser import parse_docx
from app.parsing.markdown_parser import parse_markdown
from app.parsing.models import (
    ContentItem,
    ContentItemType,
    HeadingSource,
    ParsedBlockType,
)


def save(doc: DocxDocument, path: Path) -> Path:
    doc.save(str(path))
    return path


def add_bold(doc: DocxDocument, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True


def write_md(tmp_path: Path, text: str, name: str = "doc.md") -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


# --- 1. DOCX heading inference ----------------------------------------------


def test_fully_bold_numbered_section_heading_is_inferred(tmp_path: Path) -> None:
    doc = DocxDocument()
    add_bold(doc, "4.1 Methods")
    parsed = parse_docx(save(doc, tmp_path / "numbered.docx"))

    block = parsed.blocks[0]
    assert block.block_type == ParsedBlockType.HEADING
    assert block.heading_level == 2
    assert block.heading_source == HeadingSource.INFERRED


def test_bold_prose_remains_a_normal_paragraph(tmp_path: Path) -> None:
    doc = DocxDocument()
    add_bold(doc, "Our results show a significant improvement in overall performance.")
    parsed = parse_docx(save(doc, tmp_path / "prose.docx"))

    block = parsed.blocks[0]
    assert block.block_type == ParsedBlockType.PARAGRAPH
    assert block.heading_source is None


def test_boldness_alone_is_not_a_heading(tmp_path: Path) -> None:
    doc = DocxDocument()
    add_bold(doc, "Important note")
    parsed = parse_docx(save(doc, tmp_path / "bold_only.docx"))
    assert parsed.blocks[0].block_type == ParsedBlockType.PARAGRAPH


def test_bold_year_style_number_is_not_a_heading(tmp_path: Path) -> None:
    doc = DocxDocument()
    add_bold(doc, "2024 was a productive year")
    parsed = parse_docx(save(doc, tmp_path / "year.docx"))
    assert parsed.blocks[0].block_type == ParsedBlockType.PARAGRAPH


def test_nested_numbered_headings_are_inferred(tmp_path: Path) -> None:
    doc = DocxDocument()
    add_bold(doc, "4 Introduction")
    add_bold(doc, "4.1 Background")
    add_bold(doc, "4.1.1 Details")
    doc.add_paragraph("Body text.")
    parsed = parse_docx(save(doc, tmp_path / "nested.docx"))

    assert [block.heading_level for block in parsed.blocks[:3]] == [1, 2, 3]
    assert parsed.blocks[-1].heading_path == [
        "4 Introduction",
        "4.1 Background",
        "4.1.1 Details",
    ]


def test_inferred_headings_build_hierarchy(tmp_path: Path) -> None:
    doc = DocxDocument()
    add_bold(doc, "4 Findings")
    add_bold(doc, "4.1 Result")
    doc.add_paragraph("The result was positive.")
    parsed = parse_docx(save(doc, tmp_path / "hierarchy.docx"))

    body = parsed.blocks[-1]
    assert body.heading_path == ["4 Findings", "4.1 Result"]


def test_explicit_word_heading_style_overrides_inference(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("4.1 Methods", level=1)  # explicit Heading 1
    parsed = parse_docx(save(doc, tmp_path / "explicit.docx"))

    block = parsed.blocks[0]
    assert block.heading_level == 1
    assert block.heading_source == HeadingSource.EXPLICIT


# --- 2. canonical ordered content_items -------------------------------------


def test_content_items_preserve_docx_interleaving(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell"
    doc.add_paragraph("Second paragraph.")
    parsed = parse_docx(save(doc, tmp_path / "interleave.docx"))

    assert [item.item_type for item in parsed.content_items] == [
        ContentItemType.BLOCK,
        ContentItemType.TABLE,
        ContentItemType.BLOCK,
    ]
    assert [item.sequence_index for item in parsed.content_items] == [0, 1, 2]
    assert parsed.content_items[0].block is not None
    assert parsed.content_items[0].block.text == "First paragraph."
    assert parsed.content_items[1].table is not None
    assert parsed.content_items[1].table.rows[0].cells[0].text == "Cell"
    assert parsed.content_items[2].block is not None
    assert parsed.content_items[2].block.text == "Second paragraph."


def test_content_items_cover_all_blocks_and_tables(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Paragraph.")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("Another paragraph.")
    parsed = parse_docx(save(doc, tmp_path / "cover.docx"))

    assert len(parsed.content_items) == len(parsed.blocks) + len(parsed.tables)


def test_content_items_preserve_markdown_order(tmp_path: Path) -> None:
    text = "# Title\n\nFirst para.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\nAfter table."
    parsed = parse_markdown(write_md(tmp_path, text))

    assert [item.item_type for item in parsed.content_items] == [
        ContentItemType.BLOCK,
        ContentItemType.BLOCK,
        ContentItemType.TABLE,
        ContentItemType.BLOCK,
    ]
    assert [item.sequence_index for item in parsed.content_items] == [0, 1, 2, 3]


def test_content_items_block_identity_matches_blocks_list(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Only paragraph.")
    parsed = parse_docx(save(doc, tmp_path / "identity.docx"))

    assert parsed.content_items[0].block is parsed.blocks[0]


def test_content_item_requires_matching_payload() -> None:
    with pytest.raises(ValidationError):
        ContentItem(sequence_index=0, item_type=ContentItemType.BLOCK)
    with pytest.raises(ValidationError):
        ContentItem(sequence_index=0, item_type=ContentItemType.TABLE)
