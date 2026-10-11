"""Shared helpers for the RAG-006 metadata tests.

Synthetic PDFs/DOCX files and an in-memory HTTP transport keep the suite fast and
fully offline: no test contacts Crossref or OpenAlex.
"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import pymupdf
from docx import Document as DocxDocument

from app.metadata.enrichment import EnrichmentComponents, MetadataEnricher
from app.metadata.external.cache import JsonCache
from app.metadata.external.clients import CrossrefClient, OpenAlexClient
from app.metadata.external.http import HttpClient, RateLimiter


class FakeResponse:
    """A minimal stand-in for an HTTP response object."""

    def __init__(self, status_code: int, data: Any = None, text: str | None = None) -> None:
        self.status_code = status_code
        self._data = data
        self.text = text if text is not None else json.dumps(data if data is not None else {})

    def json(self) -> Any:
        if self._data is None:
            raise ValueError("no JSON body")
        return self._data


class FakeTransport:
    """Records calls and delegates to an optional handler."""

    def __init__(self, handler: Callable[[str, dict | None], FakeResponse] | None = None) -> None:
        self.calls: list[dict] = []
        self.handler = handler

    def get(self, url: str, *, params=None, headers=None, timeout: float = 10.0) -> FakeResponse:
        self.calls.append({"url": url, "params": params, "headers": headers})
        if self.handler is not None:
            return self.handler(url, params)
        return FakeResponse(404, {})


def make_http(transport: FakeTransport, *, max_retries: int = 0) -> HttpClient:
    """Build an :class:`HttpClient` that never sleeps and does no rate limiting."""
    return HttpClient(
        transport,
        max_retries=max_retries,
        rate_limiter=RateLimiter(0.0, sleep=lambda _seconds: None),
        sleep=lambda _seconds: None,
    )


def build_test_enricher(
    transport: FakeTransport,
    *,
    cache: JsonCache | None = None,
    crossref: bool = True,
    openalex: bool = True,
    enabled: bool = True,
    max_retries: int = 0,
) -> MetadataEnricher:
    """Build an enricher wired to a fake transport."""
    http = make_http(transport, max_retries=max_retries)
    return MetadataEnricher(
        crossref=CrossrefClient(http, cache) if crossref else None,
        openalex=OpenAlexClient(http, cache) if openalex else None,
        enabled=enabled,
    )


def build_test_stack(
    transport: FakeTransport,
    cache: JsonCache,
    *,
    crossref: bool = True,
    openalex: bool = True,
    enabled: bool = True,
    title_lookup_enabled: bool = True,
    max_retries: int = 0,
) -> tuple[MetadataEnricher, EnrichmentComponents]:
    """Build an enricher *and* its components (for stage-runner tests)."""
    http = make_http(transport, max_retries=max_retries)
    crossref_client = CrossrefClient(http, cache) if crossref else None
    openalex_client = OpenAlexClient(http, cache) if openalex else None
    enricher = MetadataEnricher(
        crossref=crossref_client,
        openalex=openalex_client,
        enabled=enabled,
        title_lookup_enabled=title_lookup_enabled,
    )
    components = EnrichmentComponents(
        http=http,
        cache=cache,
        crossref=crossref_client,
        openalex=openalex_client,
    )
    return enricher, components


def crossref_message(
    *,
    doi: str = "10.1000/xyz",
    title: str = "A Verified Title",
    authors: list[tuple[str, str]] | None = None,
    year: int = 2019,
    type_: str = "journal-article",
    venue: str = "Journal of Testing",
    publisher: str = "Test Press",
) -> dict:
    """Build a Crossref ``message`` payload."""
    author_list = authors if authors is not None else [("Jane", "Doe")]
    return {
        "DOI": doi,
        "title": [title],
        "author": [
            {"given": given, "family": family} for given, family in author_list
        ],
        "issued": {"date-parts": [[year, 1, 1]]},
        "type": type_,
        "container-title": [venue],
        "publisher": publisher,
        "URL": f"https://doi.org/{doi}",
    }


def crossref_doi_handler(message: dict, *, status: int = 200) -> Callable:
    """Return a transport handler answering Crossref DOI lookups."""

    def handler(url: str, params: dict | None) -> FakeResponse:
        if "/works/" in url:
            return FakeResponse(status, {"message": message} if status == 200 else {})
        return FakeResponse(404, {})

    return handler


def crossref_search_handler(messages: list[dict]) -> Callable:
    """Return a transport handler answering Crossref bibliographic searches."""

    def handler(url: str, params: dict | None) -> FakeResponse:
        if url.rstrip("/").endswith("/works") and params and "query.bibliographic" in params:
            return FakeResponse(200, {"message": {"items": messages}})
        return FakeResponse(404, {})

    return handler


def make_pdf(
    tmp_path: Path,
    name: str = "doc.pdf",
    *,
    metadata: dict | None = None,
    lines: list[tuple[str, float, float]] | None = None,
) -> Path:
    """Create a synthetic PDF with the given metadata and (text, size, y) lines."""
    return make_pdf_pages(tmp_path, name, [lines or []], metadata=metadata)


def make_pdf_pages(
    tmp_path: Path,
    name: str,
    pages: list[list[tuple[str, float, float]]],
    *,
    metadata: dict | None = None,
) -> Path:
    """Create a synthetic multi-page PDF; ``pages`` is a list of line lists."""
    document = pymupdf.open()
    for page_lines in pages:
        page = document.new_page()
        for text, size, y in page_lines:
            page.insert_text((72, y), text, fontsize=size)
    if metadata is not None:
        document.set_metadata(metadata)
    path = tmp_path / name
    document.save(str(path))
    document.close()
    return path


def make_docx(
    tmp_path: Path,
    name: str = "doc.docx",
    *,
    title: str | None = None,
    author: str | None = None,
    created: datetime | None = None,
    paragraphs: list[str] | None = None,
    title_style: bool = False,
) -> Path:
    """Create a synthetic DOCX with the given core properties and paragraphs."""
    document = DocxDocument()
    if title_style and title:
        document.add_paragraph(title, style="Title")
    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)
    core = document.core_properties
    if title is not None:
        core.title = title
    if author is not None:
        core.author = author
    if created is not None:
        core.created = created
    path = tmp_path / name
    document.save(str(path))
    return path


def make_markdown(tmp_path: Path, name: str = "doc.md", text: str = "") -> Path:
    """Create a Markdown file with the given content."""
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def utc(year: int, month: int = 1, day: int = 1) -> datetime:
    """A timezone-aware UTC datetime for DOCX core-property tests."""
    return datetime(year, month, day, tzinfo=timezone.utc)
