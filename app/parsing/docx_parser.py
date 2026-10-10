"""DOCX parser.

Uses ``python-docx`` to preserve document order (paragraphs and tables
interleaved), Word heading styles, list items, quotes, tables, and the heading
hierarchy. Files are opened read-only; the source document is never modified.

Explicit Word heading styles are authoritative. When a document carries no
heading metadata (as the real thesis files do), a conservative deterministic
fallback infers headings from structural signals only.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsing.exceptions import DocumentParseError
from app.parsing.models import (
    ContentItem,
    ContentItemType,
    HeadingSource,
    ParsedBlock,
    ParsedBlockType,
    ParsedDocument,
    ParsedTable,
    ParsedTableCell,
    ParsedTableRow,
)

_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)
_TITLE_STYLE = "title"

#: Heading level used for a paragraph explicitly styled as the document Title.
#: It sits above Heading 1, so it roots the heading hierarchy.
_TITLE_LEVEL = 0

#: Conservative inference signals for DOCX files that carry no Word heading
#: metadata. Inference is a fallback only and never fires on boldness alone.
_CHAPTER_RE = re.compile(r"^(?:chapter|appendix)\s+\d+\b", re.IGNORECASE)
_NUMBERED_RE = re.compile(r"^(?P<num>\d+(?:\.\d+)*)(?P<rest>.*)$")
_MAX_HEADING_CHARS = 100
_MAX_HEADING_WORDS = 12
_MAX_NUMBER_COMPONENT_DIGITS = 3
_MAX_NUMBER_REMAINDER_WORDS = 8


def _iter_body_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in true document order.

    ``python-docx`` exposes ``document.paragraphs`` and ``document.tables``
    separately, which loses their interleaving. Walking the body XML children
    directly preserves the original order without extra dependencies.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _heading_level(style_name: str) -> int | None:
    """Return the numeric level for a Word heading style, else ``None``."""
    match = _HEADING_STYLE_RE.match(style_name)
    return int(match.group(1)) if match else None


def _has_numbering(paragraph: Paragraph) -> bool:
    """Return ``True`` if the paragraph carries Word numbering properties."""
    paragraph_properties = getattr(paragraph._p, "pPr", None)
    if paragraph_properties is None:
        return False
    return getattr(paragraph_properties, "numPr", None) is not None


def _is_fully_bold(paragraph: Paragraph) -> bool:
    """Return ``True`` when every non-empty run is explicitly bold."""
    text_runs = [run for run in paragraph.runs if run.text and run.text.strip()]
    if not text_runs:
        return False
    return all(run.bold for run in text_runs)


def _looks_like_prose(text: str) -> bool:
    """Return ``True`` when text reads like an ordinary sentence, not a heading."""
    if len(text) > _MAX_HEADING_CHARS:
        return True
    if len(text.split()) > _MAX_HEADING_WORDS:
        return True
    return text.endswith((".", "!", "?"))


def _infer_heading_level(paragraph: Paragraph, text: str) -> int | None:
    """Conservatively infer a heading level, or ``None`` to leave it as prose.

    A paragraph must be fully bold **and** match a strong structural pattern
    (a chapter label or a numbered section) to qualify. Boldness alone never
    makes a heading, and font sizes are not used (unreliable in this corpus).
    """
    stripped = text.strip()
    if not stripped or not _is_fully_bold(paragraph) or _looks_like_prose(stripped):
        return None

    if _CHAPTER_RE.match(stripped):
        return 1

    match = _NUMBERED_RE.match(stripped)
    if match is None:
        return None

    components = match.group("num").split(".")
    if any(len(component) > _MAX_NUMBER_COMPONENT_DIGITS for component in components):
        return None

    remainder = match.group("rest").strip().lstrip(").").strip()
    if remainder:
        if not (remainder[0].isupper() or remainder[0].isdigit()):
            return None
        if len(remainder.split()) > _MAX_NUMBER_REMAINDER_WORDS:
            return None

    return min(len(components), 6)


def _render_table(rows: list[ParsedTableRow]) -> str:
    """Render a readable plain-text view of a table for ``raw_text`` only."""
    return "\n".join(" | ".join(cell.text for cell in row.cells) for row in rows)


def parse_docx(path: Path, file_hash: str | None = None) -> ParsedDocument:
    """Parse a ``.docx`` file into a :class:`ParsedDocument`.

    Raises:
        DocumentParseError: if the file is missing, cannot be opened, or is not
            a valid DOCX package.
    """
    path = Path(path)
    if not path.is_file():
        raise DocumentParseError(f"Document not found: {path}")

    try:
        document = DocxDocument(str(path))
    except Exception as exc:  # python-docx raises various package errors
        raise DocumentParseError(f"Failed to parse DOCX '{path}': {exc}") from exc

    blocks: list[ParsedBlock] = []
    tables: list[ParsedTable] = []
    content_items: list[ContentItem] = []
    raw_parts: list[str] = []

    headings: dict[int, str] = {}
    title: str | None = None
    paragraph_counter = 0
    table_counter = 0
    block_counter = 0
    sequence_counter = 0

    def heading_path() -> list[str]:
        return [text for _level, text in sorted(headings.items())]

    def emit_block(block: ParsedBlock) -> None:
        nonlocal block_counter, sequence_counter
        blocks.append(block)
        content_items.append(
            ContentItem(
                sequence_index=sequence_counter,
                item_type=ContentItemType.BLOCK,
                block=block,
            )
        )
        block_counter += 1
        sequence_counter += 1

    for item in _iter_body_items(document):
        if isinstance(item, Paragraph):
            index = paragraph_counter
            paragraph_counter += 1

            text = item.text.strip()
            style_name = item.style.name if item.style is not None else ""
            style_name = (style_name or "").strip()

            if style_name.casefold() == _TITLE_STYLE:
                if not text:
                    continue
                if title is None:
                    title = text
                _update_headings(headings, _TITLE_LEVEL, text)
                emit_block(
                    _make_block(
                        block_counter,
                        sequence_counter,
                        ParsedBlockType.HEADING,
                        text,
                        heading_level=_TITLE_LEVEL,
                        heading_source=HeadingSource.EXPLICIT,
                        heading_path=heading_path(),
                        paragraph_index=index,
                    )
                )
                raw_parts.append(text)
                continue

            if not text:
                # Empty paragraphs carry no readable content.
                continue

            level = _heading_level(style_name)
            heading_source = HeadingSource.EXPLICIT if level is not None else None
            is_list_item = style_name.startswith("List") or _has_numbering(item)
            is_quote = style_name.startswith("Quote")

            if level is None and not is_list_item and not is_quote:
                inferred_level = _infer_heading_level(item, text)
                if inferred_level is not None:
                    level = inferred_level
                    heading_source = HeadingSource.INFERRED

            if level is not None:
                _update_headings(headings, level, text)
                emit_block(
                    _make_block(
                        block_counter,
                        sequence_counter,
                        ParsedBlockType.HEADING,
                        text,
                        heading_level=level,
                        heading_source=heading_source,
                        heading_path=heading_path(),
                        paragraph_index=index,
                    )
                )
                raw_parts.append(text)
                continue

            if is_list_item:
                block_type = ParsedBlockType.LIST_ITEM
            elif is_quote:
                block_type = ParsedBlockType.QUOTE
            else:
                block_type = ParsedBlockType.PARAGRAPH

            emit_block(
                _make_block(
                    block_counter,
                    sequence_counter,
                    block_type,
                    text,
                    heading_path=heading_path(),
                    paragraph_index=index,
                )
            )
            raw_parts.append(text)
        else:
            index = table_counter
            table_counter += 1
            rows = [
                ParsedTableRow(
                    cells=[ParsedTableCell(text=cell.text.strip()) for cell in row.cells]
                )
                for row in item.rows
            ]
            parsed_table = ParsedTable(
                sequence_index=sequence_counter,
                table_index=index,
                rows=rows,
                heading_path=heading_path(),
                source_locator=f"table:{index}",
            )
            tables.append(parsed_table)
            content_items.append(
                ContentItem(
                    sequence_index=sequence_counter,
                    item_type=ContentItemType.TABLE,
                    table=parsed_table,
                )
            )
            sequence_counter += 1
            rendered = _render_table(rows)
            if rendered:
                raw_parts.append(rendered)

    if title is None:
        core_title = document.core_properties.title
        if core_title and core_title.strip():
            title = core_title.strip()

    return ParsedDocument(
        source_path=str(path),
        filename=path.name,
        file_hash=file_hash,
        extension=".docx",
        title=title,
        blocks=blocks,
        tables=tables,
        content_items=content_items,
        raw_text="\n\n".join(raw_parts),
    )


def _update_headings(headings: dict[int, str], level: int, text: str) -> None:
    """Set ``level`` to ``text`` and drop any deeper levels."""
    for existing in [key for key in headings if key >= level]:
        del headings[existing]
    headings[level] = text


def _make_block(
    block_index: int,
    sequence_index: int,
    block_type: ParsedBlockType,
    text: str,
    *,
    heading_level: int | None = None,
    heading_source: HeadingSource | None = None,
    heading_path: list[str],
    paragraph_index: int | None = None,
) -> ParsedBlock:
    """Build a :class:`ParsedBlock` with a deterministic source locator."""
    locator = f"paragraph:{paragraph_index}" if paragraph_index is not None else None
    return ParsedBlock(
        sequence_index=sequence_index,
        block_index=block_index,
        block_type=block_type,
        text=text,
        heading_level=heading_level,
        heading_source=heading_source,
        heading_path=list(heading_path),
        paragraph_index=paragraph_index,
        source_locator=locator,
    )
